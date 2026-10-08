"""Profil ``gh_standard_v1`` (GH-Vorlage mit Geschäftsvorfall, Varianten und ohne Systemspalte)."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.shared import Inches

from tests.fixtures.docx_factory import PNG, build_gh_docx
from tests.helpers import RunConvert, file_entry, load_report, load_testcase


def test_gh_vorlage_wird_erkannt_und_exportiert(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    build_gh_docx(
        input_dir / "gh.docx",
        [
            ("1", "", "MEK1 - EK-Konditionen anlegen", "", "", ""),
            ("1.1", "0001", "Einkaufskondition anlegen", "Lieferant", "Einkaufsorganisation 0010", "Eintrag wird angelegt."),
        ],
    )

    assert run_convert() == 0

    entry = file_entry(load_report(output_dir), "gh.docx")
    assert entry["detected_profile"] == "gh_standard_v1"
    assert [check["profile"] for check in entry["checked_profiles"] if check["matched"]] == ["gh_standard_v1"]
    assert entry["warnings"] == []

    testcase = load_testcase(output_dir, "gh")
    assert testcase["summary"] == "TFB_GH_0001"
    assert testcase["labels"] == ["RWWS-GH", "Erika_Muster"]
    assert testcase["custom_fields"] == {
        "customfield_15909": "03.02 Einkaufsverwaltung/03.02.001 GH Pflege Einkaufskonditionen (EGKE)"
    }
    # Info-Tabelle ist die Tabelle mit "Geschäftsvorfall", nicht die letzte Tabelle vor dem Testablauf.
    assert testcase["description"] == (
        "h1. MEK1 - EK-Konditionen anlegen\n\n"
        "|Geschäftsvorfall|MEK1 - EK-Konditionen anlegen|Verantwortlicher Tester| |\n"
        "|Beschreibung| | | |"
    )
    first, second = testcase["steps"]
    assert first["system"] == "nicht definiert"
    assert first["action"] == "MEK1 - EK-Konditionen anlegen"
    assert first["expected_result"] == "-"
    assert second["action"] == "Einkaufskondition anlegen\nEinkaufsorganisation 0010"
    assert second["data"] == ""
    assert second["expected_result"] == "Eintrag wird angelegt."
    assert "Lieferant" not in str(testcase) and "0001" not in str(testcase["steps"])


def test_gh_eingabedaten_text_und_bilder_folgen_der_aktion(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    path = build_gh_docx(
        input_dir / "dateiname_anders.docx",
        [
            ("1", "", "Geschäftsprozess-Schritt", "", "Eingabedaten", "Ergebnis"),
            ("2", "", "Zweiter Geschäftsprozess-Schritt", "", "", "Ergebnis 2"),
        ],
        name="TFB_GH – Prüfung Öl & Sonderzeichen",
    )
    document = Document(str(path))
    table = next(
        item
        for item in document.tables
        if any(cell.text == "Geschäftsprozess-Schritte" for cell in item.rows[0].cells)
    )
    action_cell = table.cell(1, 2)
    input_cell = table.cell(1, 4)
    for cell in (action_cell, input_cell, input_cell, table.cell(2, 4)):
        cell.add_paragraph().add_run().add_picture(BytesIO(PNG), width=Inches(0.2))
    document.save(str(path))

    assert run_convert() == 0

    testcase = load_testcase(output_dir, "dateiname_anders")
    step = testcase["steps"][0]
    assert testcase["summary"] == "TFB_GH – Prüfung Öl & Sonderzeichen"
    assert step["action"] == (
        "Geschäftsprozess-Schritt\n!0001.png!\nEingabedaten\n!0002.png!\n!0003.png!"
    )
    assert step["data"] == ""
    assert step["expected_result"] == "Ergebnis"
    assert step["attachments"] == ["0001.png", "0002.png", "0003.png"]
    second_step = testcase["steps"][1]
    assert second_step["action"] == "Zweiter Geschäftsprozess-Schritt\n!0004.png!"
    assert second_step["data"] == ""
    assert second_step["expected_result"] == "Ergebnis 2"
    assert second_step["attachments"] == ["0004.png"]
    assert testcase["screenshots"] == []
    assert sorted(path.name for path in (output_dir / "dateiname_anders" / "screenshots").iterdir()) == [
        "0001.png",
        "0002.png",
        "0003.png",
        "0004.png",
    ]
