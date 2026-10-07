"""Automatische Profilerkennung und CLI-Override ``--profile``."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from tests.fixtures.docx_factory import DocSpec, StepSpec, build_lunar_docx
from tests.helpers import RunConvert, error_codes, file_entry, final_entries, load_report, load_testcase


def test_dokument_ohne_passendes_profil(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    build_lunar_docx(input_dir / "fremd.docx", DocSpec(include_kurzbeschreibung=False, include_testablauf=False))

    assert run_convert() == 1

    entry = file_entry(load_report(output_dir), "fremd.docx")
    assert entry["profile_detection_status"] == "no_matching_profile"
    assert error_codes(entry) == ["no_matching_profile"]
    assert final_entries(output_dir) == []


def test_fehlende_pflichtspalte_passt_nicht(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    header = ("Schritt-Nr.", "System / Komponente", "Beschreibung des Testschritts", "Bemerkung")
    build_lunar_docx(input_dir / "spalte.docx", DocSpec(step_header=header))

    assert run_convert() == 1

    entry = file_entry(load_report(output_dir), "spalte.docx")
    assert entry["profile_detection_status"] == "no_matching_profile"
    assert all(any("Pflichtspalten" in reason for reason in check["reasons"]) for check in entry["checked_profiles"])


def test_mehrere_passende_profile_sind_mehrdeutig(
    input_dir: Path, output_dir: Path, run_convert: RunConvert, two_profile_config: Path
) -> None:
    build_lunar_docx(input_dir / "a.docx")

    assert run_convert(config_dir=two_profile_config) == 1

    entry = file_entry(load_report(output_dir), "a.docx")
    assert entry["profile_detection_status"] == "ambiguous_profile"
    assert entry["detected_profile"] is None
    assert error_codes(entry) == ["ambiguous_profile"]
    assert [check["matched"] for check in entry["checked_profiles"]] == [True, True]
    assert final_entries(output_dir) == []


def test_vorlage_ohne_datenzeile_wird_ohne_steps_exportiert(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    empty_rows = [StepSpec(number="", system="", action="", expected=""), StepSpec(number="20", system="", action="", expected="")]
    build_lunar_docx(input_dir / "vorlage.docx", DocSpec(steps=empty_rows))

    assert run_convert() == 0

    entry = file_entry(load_report(output_dir), "vorlage.docx")
    assert entry["profile_detection_status"] == "matched"
    standard = next(check for check in entry["checked_profiles"] if check["profile"] == "lunar_standard_v1")
    assert any("Keine fachlich befüllte Schrittzeile" in reason for reason in standard["reasons"])
    testcase = load_testcase(output_dir, "vorlage")
    assert "steps" not in testcase
    assert testcase["summary"] == "TF_Beispiel_001"


def test_fico_dokument_ohne_nummerierung_und_mit_name_in_tabellenzelle(
    input_dir: Path, output_dir: Path, run_convert: RunConvert
) -> None:
    document = Document()
    document.add_paragraph("Produkt: FiCO")
    info = document.add_table(rows=5, cols=2)
    for row, (label, value) in enumerate(
        [
            ("Testfalltitel:", "TF_FICO_1"),
            ("Kurzbeschreibung:", "Kurz"),
            ("Testvoraussetzungen:", "Keine"),
            ("Erwartetes Ergebnis:", "OK"),
            ("Ergebnisprüfung:", "Sichtprüfung"),
        ]
    ):
        info.cell(row, 0).text = label
        info.cell(row, 1).text = value
    document.add_paragraph("Testablauf:")
    header = ("Lfd. Nr.", "Beschreibung des Schrittes", "Transaktion/ Report", "Erwartetes Ergebnis", "Tatsächliches Ergebnis / Fehlerbeschreibung")
    steps = document.add_table(rows=2, cols=len(header))
    for column, text in enumerate(header):
        steps.cell(0, column).text = text
    part = document.part
    for rel_id, rel in list(part.rels.items()):
        if rel.reltype == RT.NUMBERING:
            part.drop_rel(rel_id)
    document.save(str(input_dir / "fico.docx"))

    assert run_convert() == 0

    entry = file_entry(load_report(output_dir), "fico.docx")
    assert entry["detected_profile"] == "fico_standard_v1"
    testcase = load_testcase(output_dir, "fico")
    assert testcase["summary"] == "TF_FICO_1"
    assert "steps" not in testcase


def test_profile_override_prueft_nur_ein_profil(
    input_dir: Path, output_dir: Path, run_convert: RunConvert, two_profile_config: Path
) -> None:
    build_lunar_docx(input_dir / "a.docx")

    assert run_convert("--profile", "lunar_kopie_v1", config_dir=two_profile_config) == 0

    report = load_report(output_dir)
    entry = file_entry(report, "a.docx")
    assert report["profile_override"] == "lunar_kopie_v1"
    assert entry["profile_detection_status"] == "matched"
    assert entry["detected_profile"] == "lunar_kopie_v1"
    assert [check["profile"] for check in entry["checked_profiles"]] == ["lunar_kopie_v1"]


def test_profile_override_ohne_fallback(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    build_lunar_docx(input_dir / "a.docx", DocSpec(include_kurzbeschreibung=False, include_testablauf=False))

    assert run_convert("--profile", "lunar_standard_v1") == 1

    entry = file_entry(load_report(output_dir), "a.docx")
    assert error_codes(entry) == ["no_matching_profile"]
    assert len(entry["checked_profiles"]) == 1


def test_unbekanntes_profil_ist_globaler_fehler(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    build_lunar_docx(input_dir / "a.docx")

    assert run_convert("--profile", "gibt_es_nicht") == 2
    assert not output_dir.exists()


def test_profilerkennung_ignoriert_dateinamen(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    build_lunar_docx(
        input_dir / "lunar_standard_v1_Testablauf.docx",
        DocSpec(include_kurzbeschreibung=False, include_testablauf=False),
    )

    assert run_convert() == 1
    assert error_codes(file_entry(load_report(output_dir), "lunar_standard_v1_Testablauf.docx")) == ["no_matching_profile"]
