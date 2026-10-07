"""Fachliche Pflichtfelder je Testfall und Testschritt."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches

from tests.fixtures.docx_factory import PNG, DocSpec, StepSpec, build_lunar_docx
from tests.helpers import RunConvert, error_codes, file_entry, final_entries, load_report, load_testcase


def test_fehlender_testfallname_wird_mit_required_field_missing_abgelehnt(
    input_dir: Path, output_dir: Path, run_convert: RunConvert
) -> None:
    build_lunar_docx(input_dir / "ohne_name.docx", DocSpec(name=None))

    assert run_convert() == 1

    entry = file_entry(load_report(output_dir), "ohne_name.docx")
    assert entry["detected_profile"] == "lunar_standard_v1"
    assert error_codes(entry) == ["required_field_missing"]
    assert entry["errors"][0]["field"] == "summary"
    assert final_entries(output_dir) == []


def test_leerer_testfallname_wird_mit_required_field_missing_abgelehnt(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    build_lunar_docx(input_dir / "leer.docx", DocSpec(name="   "))

    assert run_convert() == 1
    entry = file_entry(load_report(output_dir), "leer.docx")
    assert error_codes(entry) == ["required_field_missing"]
    assert entry["errors"][0]["field"] == "summary"


def test_mehrere_unterschiedliche_namensfelder_werden_abgelehnt(
    input_dir: Path, output_dir: Path, run_convert: RunConvert
) -> None:
    path = build_lunar_docx(input_dir / "mehrere_namen.docx")
    document = Document(str(path))
    document.tables[0].cell(1, 0).add_paragraph("Testfallname: Anderer Name")
    document.save(str(path))

    assert run_convert() == 1

    entry = file_entry(load_report(output_dir), "mehrere_namen.docx")
    assert error_codes(entry) == ["required_field_missing"]
    assert any(warning["code"] == "ambiguous_testcase_name" for warning in entry["warnings"])


def test_legacy_formcheckbox_wird_vor_textverarbeitung_gerendert(
    input_dir: Path, output_dir: Path, run_convert: RunConvert
) -> None:
    path = build_lunar_docx(input_dir / "legacy_checkbox.docx")
    document = Document(str(path))
    title_paragraph = document.tables[0].cell(1, 0).paragraphs[0]
    title_paragraph._p.append(
        parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:fldChar w:fldCharType="begin"><w:ffData><w:checkBox><w:default w:val="1"/>'
            "</w:checkBox></w:ffData></w:fldChar></w:r>"
        )
    )
    title_paragraph.add_run(" Testsystem")

    info_table = next(table for table in document.tables if any(cell.text == "Kurzbeschreibung" for row in table.rows for cell in row.cells))
    paragraph = info_table.cell(1, 1).paragraphs[0]
    for value, checked in (("An", "1"), ("Aus", "0")):
        paragraph._p.append(
            parse_xml(
                '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:fldChar w:fldCharType="begin"><w:ffData><w:checkBox>'
                f'<w:default w:val="{checked}"/>'
                "</w:checkBox></w:ffData></w:fldChar></w:r>"
            )
        )
        paragraph.add_run(f" {value} ")
    paragraph._p.append(
        parse_xml(
            '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            '<w:fldChar w:fldCharType="begin"><w:ffData><w:textInput/></w:ffData></w:fldChar>'
            "</w:r>"
        )
    )
    paragraph.add_run(" FORMTEXT")
    step_table = next(
        table for table in document.tables
        if any(cell.text == "Beschreibung des Testschritts" for cell in table.rows[0].cells)
    )
    unknown_glyph = step_table.cell(1, 2).add_paragraph().add_run("\uf0b0")
    unknown_glyph.font.name = "Wingdings"
    document.save(str(path))

    assert run_convert() == 0

    testcase = load_testcase(output_dir, "legacy_checkbox")
    assert "(/) An (x) Aus FORMTEXT" in testcase["description"]
    assert "[Wingdings U+F0B0]" in testcase["steps"][0]["action"]
    entry = file_entry(load_report(output_dir), "legacy_checkbox.docx")
    assert entry["errors"] == []
    assert not any("CheckboxRef" in warning["message"] for warning in entry["warnings"])
    wingdings_warning = next(warning for warning in entry["warnings"] if warning["code"] == "unsupported_wingdings_glyph")
    assert "legacy_checkbox.docx" in wingdings_warning["message"]
    assert "Tabelle" in wingdings_warning["message"] and "U+F0B0" in wingdings_warning["message"]


def test_teilweise_befuellte_schritte_und_leere_zeilen(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    steps = [
        StepSpec(number="5", action="", expected="", system="SAP FI", actual="Nur tatsächlich"),
        StepSpec(number="10", action="Nur Aktion", expected="", system=""),
        StepSpec(number="20", action="", expected="Nur erwartetes Ergebnis", system="SAP SD"),
        StepSpec(number="30", action="", expected="", system=""),
    ]
    build_lunar_docx(input_dir / "teilweise.docx", DocSpec(steps=steps))

    assert run_convert() == 0

    testcase = load_testcase(output_dir, "teilweise")
    assert len(testcase["steps"]) == 2
    assert testcase["steps"][0]["action"] == "Nur Aktion"
    assert testcase["steps"][0]["expected_result"] == "-"
    assert testcase["steps"][0]["system"] == "nicht definiert"
    assert testcase["steps"][1]["action"] == "-"
    assert testcase["steps"][1]["expected_result"] == "Nur erwartetes Ergebnis"
    assert testcase["steps"][1]["system"] == "SAP SD"


def test_leeres_erwartetes_ergebnis_wird_zu_strich(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    steps = [StepSpec(expected=""), StepSpec(number="20", expected="", actual="Nur tatsächlich")]
    path = build_lunar_docx(input_dir / "a.docx", DocSpec(steps=steps))
    document = Document(str(path))
    step_table = next(
        table for table in document.tables
        if any(cell.text == "Tatsächliche Ergebnisse" for cell in table.rows[0].cells)
    )
    step_table.cell(2, 4).paragraphs[0].add_run().add_picture(BytesIO(PNG), width=Inches(0.2))
    step_table.add_row()
    ignored_row = step_table.rows[-1]
    ignored_row.cells[0].text = "30"
    ignored_row.cells[1].text = "SAP FI"
    ignored_row.cells[2].text = ""
    ignored_row.cells[3].text = ""
    ignored_row.cells[4].text = "Nur tatsächliches Ergebnis"
    ignored_row.cells[4].paragraphs[0].add_run().add_picture(BytesIO(PNG), width=Inches(0.2))
    document.save(str(path))

    assert run_convert() == 0

    testcase = load_testcase(output_dir, "a")
    assert testcase["steps"][0]["expected_result"] == "-"
    assert testcase["steps"][1]["expected_result"] == "-"
    assert testcase["screenshots"] == []
    assert testcase["steps"][1]["attachments"] == []
    assert file_entry(load_report(output_dir), "a.docx")["exported_image_count"] == 0
    assert file_entry(load_report(output_dir), "a.docx")["warnings"] == []


def test_leeres_system_wird_nicht_definiert(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    steps = [StepSpec(), StepSpec(number="", system="", action="Folgeschritt", expected="Ok")]
    build_lunar_docx(input_dir / "a.docx", DocSpec(steps=steps))

    assert run_convert() == 0

    testcase = load_testcase(output_dir, "a")
    assert [step["system"] for step in testcase["steps"]] == ["SAP FI", "nicht definiert"]


def test_alle_fehlenden_felder_werden_gemeldet(input_dir: Path, output_dir: Path, run_convert: RunConvert) -> None:
    build_lunar_docx(input_dir / "a.docx", DocSpec(name=None, steps=[StepSpec(system="", action="", expected="Ergebnis")]))

    assert run_convert() == 1

    fields = [error["field"] for error in file_entry(load_report(output_dir), "a.docx")["errors"]]
    assert fields == ["summary"]
