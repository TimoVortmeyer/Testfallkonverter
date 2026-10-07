"""Profil-Preflight mit CSV-Ausgabe."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from lunar_converter import preflight
from lunar_converter.cli import main
from tests.fixtures.docx_factory import DocSpec, build_lunar_docx
from tests.helpers import CONFIG_DIR


def test_preflight_schreibt_profiltreffer_und_pruefgruende(tmp_path: Path, capsys) -> None:
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    build_lunar_docx(input_dir / "treffer.docx")
    build_lunar_docx(input_dir / "kein_treffer.docx", DocSpec(include_kurzbeschreibung=False, include_testablauf=False))
    csv_path = tmp_path / "reports" / "profile.csv"

    result = main(["preflight", "--input-dir", str(input_dir), "--csv-path", str(csv_path), "--config-dir", str(CONFIG_DIR)])

    assert result == 1
    terminal = capsys.readouterr().out
    assert "Profil-Preflight [############------------] 1/2 (50%)" in terminal
    assert "Profil-Preflight [########################] 2/2 (100%)" in terminal
    assert "Gesamt ~" in terminal and "Rest ~" in terminal
    with csv_path.open(encoding="utf-8-sig", newline="") as source:
        rows = {row["datei"]: row for row in csv.DictReader(source, delimiter=";")}
    assert rows["treffer.docx"]["status"] == "matched"
    assert rows["treffer.docx"]["erkanntes_profil"] == "lunar_standard_v1"
    assert rows["kein_treffer.docx"]["status"] == "no_matching_profile"
    assert "Marker fehlen" in rows["kein_treffer.docx"]["profilpruefungen"]
    log_path = csv_path.with_name(f"{csv_path.stem}.preflight.log")
    log_content = log_path.read_text(encoding="utf-8")
    assert "Start des Profil-Preflights" in log_content
    assert "treffer.docx" in log_content and "lunar_standard_v1" in log_content
    assert "Kein Profil passt" in log_content
    assert "Ende des Profil-Preflights" in log_content


def test_preflight_wertet_verbundene_tabellenzellen_aus(tmp_path: Path) -> None:
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    path = build_lunar_docx(input_dir / "verbunden.docx")
    document = Document(str(path))
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))
    trailing = table.rows[1]._tr
    trailing.get_or_add_trPr().append(parse_xml(f'<w:gridBefore {nsdecls("w")} w:val="1"/>'))
    document.save(str(path))
    csv_path = tmp_path / "profile.csv"

    result = main(["preflight", "--input-dir", str(input_dir), "--csv-path", str(csv_path), "--config-dir", str(CONFIG_DIR)])

    assert result == 0
    with csv_path.open(encoding="utf-8-sig", newline="") as source:
        row = next(csv.DictReader(source, delimiter=";"))
    assert row["status"] == "matched", row["fehler"]


def test_preflight_protokolliert_stacktrace_und_aufschluesselung(tmp_path: Path, monkeypatch) -> None:
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    build_lunar_docx(input_dir / "a_treffer.docx")
    build_lunar_docx(input_dir / "b_defekt.docx")
    original_read = preflight.read_docx

    def flaky_read(path: Path):
        if path.name == "b_defekt.docx":
            raise RuntimeError("Testfehler")
        return original_read(path)

    monkeypatch.setattr(preflight, "read_docx", flaky_read)
    csv_path = tmp_path / "profile.csv"

    result = main(
        ["preflight", "--input-dir", str(input_dir), "--csv-path", str(csv_path), "--config-dir", str(CONFIG_DIR), "--log-level", "DEBUG"]
    )

    assert result == 1
    log_content = csv_path.with_name(f"{csv_path.stem}.preflight.log").read_text(encoding="utf-8")
    assert "Traceback (most recent call last)" in log_content
    assert "RuntimeError: Testfehler" in log_content
    assert "1 eindeutig erkannt" in log_content
    assert "0 Eingabefehler, 1 unerwartete Fehler" in log_content


def test_preflight_konvertiert_doc_nur_temporär(tmp_path: Path, monkeypatch, capsys) -> None:
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    source_doc = input_dir / "beispiel.doc"
    source_doc.write_bytes(b"fake doc")
    template = build_lunar_docx(tmp_path / "template.docx")
    converted: list[Path] = []

    def fake_convert(source: Path, work_dir: Path) -> Path:
        converted.append(source)
        output = work_dir / "fake.docx"
        output.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(template, output)
        return output

    monkeypatch.setattr(preflight, "convert_doc_to_docx", fake_convert)
    csv_path = tmp_path / "profile.csv"

    result = main(["preflight", "--input-dir", str(input_dir), "--csv-path", str(csv_path), "--config-dir", str(CONFIG_DIR)])

    assert result == 0
    assert "Konvertiere DOC nach DOCX" in capsys.readouterr().out
    assert converted == [source_doc]
    assert source_doc.read_bytes() == b"fake doc"
    with csv_path.open(encoding="utf-8-sig", newline="") as source:
        row = next(csv.DictReader(source, delimiter=";"))
    assert row["status"] == "matched"
    assert row["erkanntes_profil"] == "lunar_standard_v1"