"""Programmatisch erzeugte DOCX-Fixtures.

Strategie:
* Es werden keine echten (ggf. vertraulichen) Kundendokumente in das Repository
  kopiert. Stattdessen erzeugt ``build_lunar_docx`` kleine DOCX-Dateien, die die
  Struktur echter LUNAR-Testdokumente nachbilden (Titeltabelle mit
  ``Testfall:``/``Verantwortlicher:``, Feldtabelle Bezeichner/Wert, Überschrift
  ``Testablauf``, Schritttabelle mit fünf Spalten).
* ``build_poc_sample_docx`` bildet exakt das Beispiel aus
  ``scripts/create_sample_doc.py`` des POC nach.
* ``build_legacy_docx`` bildet die ältere Vorlage (Profil ``lunar_legacy_v1``)
  nach: keine Systemspalte, Spalten „Feld“ und „Eingabedaten“.
* Bilder werden mit Pillow im Speicher erzeugt (PNG, JPEG, GIF als nicht
  unterstütztes Format).
* Liegt ``input/sample.docx`` lokal vor, wird es von einem optionalen
  Integrationstest zusätzlich genutzt (sonst übersprungen).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Inches
from docx.table import _Cell
from PIL import Image

STEP_HEADER: tuple[str, ...] = (
    "Schritt-Nr.",
    "System / Komponente",
    "Beschreibung des Testschritts",
    "Erwartete Ergebnisse",
    "Tatsächliche Ergebnisse",
)
LEGACY_STEP_HEADER: tuple[str, ...] = (
    "Schritt-Nr.",
    "Beschreibung des Testschritts",
    "Feld",
    "Eingabedaten\n(ggf. besondere Angaben)",
    "Erwartete Ergebnisse /\nAusgabedaten",
)
GH_STEP_HEADER: tuple[str, ...] = (
    "Schritt Nr.",
    "Variante",
    "Geschäftsprozess-Schritte",
    "Feld",
    "Eingabedaten / besondere Angaben",
    "Ausgabedaten / erwartetes Ergebnis",
)


def image_bytes(image_format: str = "PNG", color: tuple[int, int, int] = (200, 30, 30), size: tuple[int, int] = (8, 6)) -> bytes:
    buffer = BytesIO()
    Image.new("RGB", size, color).save(buffer, format=image_format)
    return buffer.getvalue()


PNG = image_bytes("PNG")
JPEG = image_bytes("JPEG", color=(30, 30, 200))
GIF = image_bytes("GIF", color=(30, 200, 30))


@dataclass
class StepSpec:
    number: str = "10"
    system: str = "SAP FI"
    action: str = "Transaktion aufrufen."
    expected: str = "Maske wird angezeigt."
    actual: str = ""
    action_images: list[bytes] = field(default_factory=list)
    expected_images: list[bytes] = field(default_factory=list)
    system_images: list[bytes] = field(default_factory=list)
    action_floating_vml_image: bytes | None = None


@dataclass
class DocSpec:
    name: str | None = "TFB_Beispiel_001"
    process_lines: tuple[str, ...] = ("03.02 Einkaufsverwaltung", "03.02.001 Pflege Einkaufskonditionen")
    title_lines: tuple[str, ...] = ("Prüfung der", "Beispielkonditionen")
    responsible: str = "Max Mustermann"
    fachbereich: str = "Finanzen"
    kurzbeschreibung: str = "Es wird ein Beispiel geprüft."
    kurzbeschreibung_images: list[bytes] = field(default_factory=list)
    voraussetzungen: str = "Stammdaten sind vorhanden."
    include_kurzbeschreibung: bool = True
    include_testablauf: bool = True
    step_header: tuple[str, ...] = STEP_HEADER
    steps: list[StepSpec] = field(default_factory=lambda: [StepSpec()])
    body_images_after_table: list[bytes] = field(default_factory=list)
    inline_vml_image_in_action: bytes | None = None
    header_images: list[bytes] = field(default_factory=list)
    footer_images: list[bytes] = field(default_factory=list)


def _add_picture(cell: _Cell, data: bytes, *, below_text: bool) -> None:
    paragraph = cell.add_paragraph() if below_text or cell.paragraphs[0].text else cell.paragraphs[0]
    paragraph.add_run().add_picture(BytesIO(data), width=Inches(0.3))


def _set_text(cell: _Cell, text: str) -> None:
    lines = text.split("\n")
    cell.paragraphs[0].text = lines[0]
    for line in lines[1:]:
        cell.add_paragraph(line)


def _add_vml_image(document: Document, cell: _Cell, data: bytes, *, floating: bool, text_before: str) -> None:
    rel_id, _ = document.part.get_or_add_image(BytesIO(data))
    style = "position:absolute;margin-left:10pt;width:20pt;height:15pt" if floating else "width:20pt;height:15pt"
    xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<w:pict><v:shape style="{style}"><v:imagedata r:id="{rel_id}"/></v:shape></w:pict></w:r>'
    )
    paragraph = cell.add_paragraph(text_before)
    paragraph._p.append(parse_xml(xml))
    paragraph.add_run(" ein.")


def build_lunar_docx(path: Path, spec: DocSpec | None = None) -> Path:
    spec = spec or DocSpec()
    document = Document()
    section = document.sections[0]
    for data in spec.header_images:
        section.header.paragraphs[0].add_run("Logo ").add_picture(BytesIO(data), width=Inches(0.3))
    for data in spec.footer_images:
        section.footer.paragraphs[0].add_run("Seite 1 ").add_picture(BytesIO(data), width=Inches(0.3))

    title = document.add_table(rows=2, cols=1)
    title.cell(0, 0).text = "RWWS"
    title_cell = title.cell(1, 0)
    cover_lines = [*spec.process_lines, *spec.title_lines]
    if spec.name is not None:
        cover_lines.append(f"Testfall: {spec.name}")
    cover_lines.append(f"Verantwortlicher: {spec.responsible}")
    _set_text(title_cell, "\n".join(cover_lines))

    document.add_paragraph("Wahrung der Vertraulichkeit")

    field_rows: list[tuple[str, str]] = [("Fachbereich", spec.fachbereich)]
    if spec.include_kurzbeschreibung:
        field_rows.append(("Kurzbeschreibung", spec.kurzbeschreibung))
    field_rows.append(("Voraussetzungen", spec.voraussetzungen))
    fields = document.add_table(rows=len(field_rows), cols=2)
    for index, (label, value) in enumerate(field_rows):
        fields.cell(index, 0).text = label
        fields.cell(index, 1).text = value
        if label == "Kurzbeschreibung":
            for data in spec.kurzbeschreibung_images:
                _add_picture(fields.cell(index, 1), data, below_text=True)

    if spec.include_testablauf:
        document.add_heading("Testablauf", level=3)

    table = document.add_table(rows=1 + len(spec.steps), cols=len(spec.step_header))
    for column, header in enumerate(spec.step_header):
        table.cell(0, column).text = header
    for row_index, step in enumerate(spec.steps, start=1):
        values = [step.number, step.system, step.action, step.expected, step.actual][: len(spec.step_header)]
        for column, value in enumerate(values):
            _set_text(table.cell(row_index, column), value)
        for data in step.system_images:
            _add_picture(table.cell(row_index, 1), data, below_text=True)
        for data in step.action_images:
            _add_picture(table.cell(row_index, 2), data, below_text=True)
        if len(spec.step_header) > 3:
            for data in step.expected_images:
                _add_picture(table.cell(row_index, 3), data, below_text=True)
        if step.action_floating_vml_image is not None:
            _add_vml_image(document, table.cell(row_index, 2), step.action_floating_vml_image, floating=True, text_before="Schwebend")
    if spec.inline_vml_image_in_action is not None and spec.steps:
        _add_vml_image(document, table.cell(1, 2), spec.inline_vml_image_in_action, floating=False, text_before="Klicken Sie auf")

    for data in spec.body_images_after_table:
        document.add_paragraph().add_run().add_picture(BytesIO(data), width=Inches(0.3))

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def build_legacy_docx(
    path: Path,
    steps: list[tuple[str, str, str, str, str]],
    name: str = "EG930_Beispiel",
    cover_label: str = "EG930_RWWS2.0_EH",
) -> Path:
    """Ältere Vorlage; ``steps`` enthält je Zeile (Nr., Beschreibung, Feld, Eingabedaten, Erwartete Ergebnisse)."""
    document = Document()
    title = document.add_table(rows=2, cols=1)
    _set_text(title.cell(0, 0), f"{cover_label}\nTestfallbeschreibung")
    _set_text(
        title.cell(1, 0),
        f"07.04 Sachkontenpflege\n07.04.001 Sachkontenstammdatenpflege\nInitiale Übernahme der Sachkonten\nTestfall: {name}",
    )

    contact = document.add_table(rows=2, cols=2)
    contact.cell(0, 0).text = "Status"
    contact.cell(0, 1).text = "Version"
    contact.cell(1, 0).text = "in Bearbeitung"
    contact.cell(1, 1).text = "1.0"

    fields = document.add_table(rows=2, cols=4)
    for column, value in enumerate(["Beschreibung", "Prüfung der Übernahme", "Datum der letzten Änderung", "31.01.2011"]):
        fields.cell(0, column).text = value
    fields.cell(1, 0).text = "Erwartete Ergebnisse"
    fields.cell(1, 1).text = "Alle Konten wurden übernommen."

    document.add_heading("Testablauf", level=3)
    table = document.add_table(rows=1 + len(steps), cols=len(LEGACY_STEP_HEADER))
    for column, header in enumerate(LEGACY_STEP_HEADER):
        _set_text(table.cell(0, column), header)
    for row_index, values in enumerate(steps, start=1):
        for column, value in enumerate(values):
            _set_text(table.cell(row_index, column), value)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def build_gh_docx(path: Path, steps: list[tuple[str, str, str, str, str, str]], name: str = "TFB_GH_0001") -> Path:
    """GH-Vorlage; ``steps`` je Zeile (Nr., Variante, Geschäftsprozess-Schritt, Feld, Eingabedaten, Ausgabedaten)."""
    document = Document()
    title = document.add_table(rows=2, cols=1)
    title.cell(0, 0).text = "RWWS-GH"
    _set_text(
        title.cell(1, 0),
        f"03.02 Einkaufsverwaltung\n03.02.001 GH Pflege Einkaufskonditionen (EGKE)\n\nMEK1 - EK-Konditionen anlegen\n"
        f"Testfall: {name}\nVerantwortlicher: Erika Muster",
    )

    info = document.add_table(rows=2, cols=4)
    for column, value in enumerate(["Geschäftsvorfall", "MEK1 - EK-Konditionen anlegen", "Verantwortlicher Tester", ""]):
        info.cell(0, column).text = value
    info.cell(1, 0).text = "Beschreibung"

    document.add_heading("Übersicht Varianten", level=3)
    variants = document.add_table(rows=2, cols=3)
    for column, value in enumerate(["Varianten-Nr", "Bezeichnung", "Beschreibung / Besonderheit"]):
        variants.cell(0, column).text = value
    variants.cell(1, 0).text = "0001"
    variants.cell(1, 1).text = "GH-Minden"

    document.add_heading("Fehlerbeschreibung / Anmerkung:", level=3)
    errors = document.add_table(rows=2, cols=2)
    errors.cell(0, 0).text = "Testablauf / Schritt Nr."
    errors.cell(0, 1).text = "Beschreibung / Anmerkungen"

    document.add_heading("Testablauf", level=3)
    table = document.add_table(rows=1 + len(steps), cols=len(GH_STEP_HEADER))
    for column, header in enumerate(GH_STEP_HEADER):
        table.cell(0, column).text = header
    for row_index, values in enumerate(steps, start=1):
        for column, value in enumerate(values):
            _set_text(table.cell(row_index, column), value)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def build_poc_sample_docx(path: Path) -> Path:
    """Nachbau von ``scripts/create_sample_doc.py`` (POC-Beispiel)."""
    document = Document()
    document.add_paragraph("Testablauf")
    document.add_paragraph("Kurzbeschreibung")
    document.add_paragraph("Testfallname: Beispiel Testfall")
    document.add_paragraph("Fachbereich: XYZ")
    rows = [
        ["Schritt-Nr.", "System / Komponente", "Beschreibung des Testschritts", "Erwartete Ergebnisse"],
        ["1", "System A", "Artikel erfassen", "Maske erscheint"],
    ]
    table = document.add_table(rows=len(rows), cols=4)
    for row_index, row_data in enumerate(rows):
        for col_index, value in enumerate(row_data):
            table.cell(row_index, col_index).text = value
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path
