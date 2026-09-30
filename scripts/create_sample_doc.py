from __future__ import annotations

from pathlib import Path

from docx import Document


def main() -> None:
    base = Path(__file__).resolve().parents[1]
    file_path = base / "input" / "sample.docx"
    file_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    document.add_paragraph("Testablauf")
    document.add_paragraph("Kurzbeschreibung")
    document.add_paragraph("Testfallname: Beispiel Testfall")
    document.add_paragraph("Fachbereich: XYZ")

    rows = [
        ["Schritt-Nr.", "System / Komponente", "Beschreibung des Testschritts", "Erwartete Ergebnisse"],
        ["1", "System A", "Artikel erfassen", "Maske erscheint"],
    ]
    table = document.add_table(rows=len(rows), cols=4)
    for row_index, row_data in enumerate(rows):
        for col_index, value in enumerate(row_data):
            table.cell(row_index, col_index).text = str(value)
    document.save(file_path)
    print(f"Erstellt: {file_path}")


if __name__ == "__main__":
    main()
